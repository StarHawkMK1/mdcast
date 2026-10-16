from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from typing import Any

from mdcast import hwpx
from mdcast.fields import RESERVED, keys_in
from mdcast.fonts import missing_primary_fonts
from mdcast.report import Codes, Report
from mdcast.templates.resolve import ResolvedTemplate, resolve_template
from mdcast.templates.schema import load_builtin_default
from mdcast.types import Format

HWPX_HEAD = "http://www.hancom.co.kr/hwpml/2011/head"

SUPPORTED_TARGETS = ["docx", "hwpx", "pdf"]


def _paragraph_texts(container: Any) -> Iterator[str]:
    for paragraph in container.paragraphs:
        yield paragraph.text
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _paragraph_texts(cell)


def _docx_parts(path: Path) -> tuple[list[str], set[str]]:
    """Header/footer paragraph texts (own parts only) and style names of a reference docx."""
    import docx

    document = docx.Document(str(path))
    texts: list[str] = []
    for section in document.sections:
        for part in (
            section.header, section.first_page_header, section.even_page_header,
            section.footer, section.first_page_footer, section.even_page_footer,
        ):
            if part.is_linked_to_previous:
                continue
            texts.extend(_paragraph_texts(part))
    return texts, {s.name for s in document.styles}


def describe_template(tpl: ResolvedTemplate) -> dict[str, Any]:
    config = tpl.config
    keys: list[str] = []
    uses_reference = False
    for part in ("header", "footer"):
        spec = config.get(part) or {}
        if spec.get("source") == "reference":
            uses_reference = True
            continue
        for slot in ("left", "center", "right"):
            keys.extend(keys_in(str(spec.get(slot) or "")))
    docx_container = tpl.container_for("docx")
    if uses_reference and docx_container.path is not None:
        try:
            texts, _ = _docx_parts(docx_container.path)
        except Exception:  # noqa: BLE001 - describe is informational; a corrupt docx adds no keys
            texts = []
        for text in texts:
            keys.extend(keys_in(text))
    defaults = (config.get("fields") or {}).get("defaults") or {}
    keys.extend(str(k) for k in defaults)
    formats = {}
    for fmt in ("docx", "hwpx", "pdf"):
        container = tpl.container_for(fmt)
        formats[fmt] = {
            "container": container.source,
            "path": str(container.path) if container.path else None,
        }
    return {
        "name": tpl.name,
        "version": tpl.version,
        "source": tpl.source,
        "location": str(tpl.location) if tpl.location else None,
        "fields": [k for k in dict.fromkeys(keys) if k not in RESERVED],
        "reserved_fields": list(RESERVED),
        "defaults": dict(defaults),
        "formats": formats,
        "header_source": (config.get("header") or {}).get("source"),
        "footer_source": (config.get("footer") or {}).get("source"),
        "table_variants": sorted(str(v) for v in (config.get("tables") or {}).get("variants") or {}),
        "supported_targets": list(SUPPORTED_TARGETS),
    }


def _check_docx_styles(tpl: ResolvedTemplate, report: Report) -> None:
    container = tpl.container_for("docx")
    if container.path is None or container.source != "custom":
        return
    try:
        _, names = _docx_parts(container.path)
    except Exception as exc:  # noqa: BLE001 - validation never rejects
        report.warn(Codes.TEMPLATE_INVALID, f"reference.docx 를 열 수 없습니다: {exc}", format="docx")
        return
    from mdcast.writers.docx.styles import PANDOC_STYLES  # lazy: writers import templates

    defaults = load_builtin_default()["styles"]
    for element, mapping in (tpl.config.get("styles") or {}).items():
        name = (mapping or {}).get("docx")
        if not name or name in names:
            continue
        if name == (defaults.get(element) or {}).get("docx") or name == (PANDOC_STYLES.get(element) or [None])[0]:
            continue
        report.warn(
            Codes.STYLE_UNMAPPED,
            f"docx 기준 문서에 스타일 '{name}' ({element}) 이(가) 없습니다 — 변환 시 기본 스타일을 씁니다",
            format="docx",
            element=element,
        )


def _check_hwpx_styles(tpl: ResolvedTemplate, report: Report) -> None:
    container = tpl.container_for("hwpx")
    if container.source != "custom" or container.path is None:
        return
    try:
        package = hwpx.Package.read(container.path.read_bytes())
        header = hwpx.Header(package.xml("Contents/header.xml"))
        names = {style.get("name") for style in header.root.iter(f"{{{HWPX_HEAD}}}style")}
    except (hwpx.ContainerError, OSError) as exc:
        report.warn(Codes.TEMPLATE_INVALID, f"reference.hwpx 를 읽을 수 없습니다: {exc}", format="hwpx")
        return
    for element, mapping in (tpl.config.get("styles") or {}).items():
        name = (mapping or {}).get("hwpx") if isinstance(mapping, dict) else None
        if name and name not in names:
            report.warn(
                Codes.STYLE_UNMAPPED,
                f"hwpx 기준 문서에 스타일 '{name}' ({element}) 이(가) 없습니다 — 변환 시 바탕글을 씁니다",
                format="hwpx",
                element=element,
            )


def validate_template(ref: str, *, cwd: Path, home: Path) -> tuple[ResolvedTemplate, Report]:
    report = Report(str(ref), Format.MD)
    tpl = resolve_template(str(ref), cwd=cwd, home=home, report=report)
    for fmt in ("docx", "hwpx"):
        tpl.note_container_fallbacks(fmt, report)
    _check_docx_styles(tpl, report)
    _check_hwpx_styles(tpl, report)
    try:
        missing = missing_primary_fonts(tpl.config.get("fonts") or {})
    except Exception as exc:  # noqa: BLE001 - font probing must not break validation
        report.warn(Codes.FONT_FALLBACK, f"글꼴 검사를 할 수 없습니다: {exc}", format="pdf")
        missing = []
    for family in missing:
        report.warn(
            Codes.FONT_FALLBACK,
            f"글꼴 '{family}' 을(를) 찾을 수 없습니다 — PDF 는 대체 글꼴로 조판됩니다",
            format="pdf",
            font=family,
        )
    report.template = tpl.info(["docx", "hwpx", "pdf"])
    return tpl, report
