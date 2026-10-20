from __future__ import annotations

import io
import struct
import subprocess
import textwrap
import zipfile
import zlib
from pathlib import Path

import docx
import pypandoc
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree

from mdcast.fields import resolve_fields
from mdcast.pipeline import build_generator, prepare_target
from mdcast.readers.md import read_md
from mdcast.report import Report
from mdcast.templates import resolve_template
from mdcast.types import Format, Options, Source


def write_template(
    root: Path, name: str, yaml_text: str, files: dict[str, bytes] | None = None
) -> Path:
    directory = root / name
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "template.yaml").write_text(textwrap.dedent(yaml_text), encoding="utf-8")
    for filename, data in (files or {}).items():
        (directory / filename).write_bytes(data)
    return directory


def tiny_png() -> bytes:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(b"\x00\xff\xff\xff")) + chunk(b"IEND", b"")


def tiny_jpeg() -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (2, 2), (200, 30, 30)).save(buffer, format="JPEG")
    return buffer.getvalue()


def tiny_bmp() -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (2, 2), (30, 30, 200)).save(buffer, format="BMP")
    return buffer.getvalue()


def build(tmp_path: Path, md: str, target: Format, template: str | None = None, *, offline: bool = True):
    report = Report("doc.md", Format.MD)
    source = Source(textwrap.dedent(md).encode("utf-8"), "doc.md", Format.MD, base_dir=tmp_path)
    document = read_md(source, report)
    ref = template or document.frontmatter.get("template")
    tpl = resolve_template(ref, cwd=tmp_path, home=tmp_path / "home", report=report)
    fields = resolve_fields(document.frontmatter, tpl.config, report)
    return prepare_target(
        document,
        target,
        tpl,
        fields,
        report,
        Options(offline=offline, cwd=tmp_path, home=tmp_path / "home"),
        tmp_path,
        tmp_path / f"work-{target.value}",
        build_generator(tpl),
    )


def make_reference_docx(
    path: Path,
    *,
    header_runs: list[str] | None = None,
    footer_runs: list[str] | None = None,
    extra_styles: tuple[str, ...] = (),
) -> Path:
    subprocess.run(
        [pypandoc.get_pandoc_path(), "-o", str(path), "--print-default-data-file", "reference.docx"],
        check=True,
    )
    document = docx.Document(str(path))
    for name in extra_styles:
        document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    section = document.sections[0]
    for part, runs in ((section.header, header_runs), (section.footer, footer_runs)):
        if runs is None:
            continue
        part.is_linked_to_previous = False
        paragraph = part.paragraphs[0]
        paragraph.text = ""
        for text in runs:
            paragraph.add_run(text)
    document.save(str(path))
    return path


def line_pitch(data: bytes, needle: str) -> float:
    """Distance (pt) between the baselines of the first two text lines from ``needle`` on page 1.

    Baselines, not character boxes: a box's top follows the glyph's outline, so lines starting
    with different syllables would seem up to ~1pt apart. The text is read per character index
    (``get_text_range`` can drift from those indices, e.g. around math).
    """
    import ctypes

    import pypdfium2 as pdfium
    import pypdfium2.raw as pdfium_c

    textpage = pdfium.PdfDocument(data)[0].get_textpage()
    count = pdfium_c.FPDFText_CountChars(textpage.raw)
    text = "".join(chr(pdfium_c.FPDFText_GetUnicode(textpage.raw, i) or 0xFFFD) for i in range(count))
    x, y = ctypes.c_double(), ctypes.c_double()
    baselines: list[float] = []
    for i in range(text.index(needle), len(text)):
        if not "가" <= text[i] <= "힣":  # body text only: a superscript sits off the baseline
            continue
        pdfium_c.FPDFText_GetCharOrigin(textpage.raw, i, ctypes.byref(x), ctypes.byref(y))
        if not baselines or abs(y.value - baselines[-1]) > 2:
            baselines.append(y.value)
        if len(baselines) == 2:
            return round(baselines[0] - baselines[1], 2)
    raise AssertionError(f"{needle!r} does not wrap")


HWPX_NS = {
    "hh": "http://www.hancom.co.kr/hwpml/2011/head",
    "hp": "http://www.hancom.co.kr/hwpml/2011/paragraph",
    "hs": "http://www.hancom.co.kr/hwpml/2011/section",
    "hc": "http://www.hancom.co.kr/hwpml/2011/core",
    "opf": "http://www.idpf.org/2007/opf/",
}
BUILTIN_HWPX = Path(__file__).resolve().parents[1] / "src/mdcast/builtin_templates/default/reference.hwpx"


def builtin_hwpx() -> bytes:
    return BUILTIN_HWPX.read_bytes()


def hwpx_integrity(data: bytes) -> dict:
    """Parse a hwpx and assert what Hangul relies on: package order, counts, unique ids, references."""
    hh = lambda tag: f"{{{HWPX_NS['hh']}}}{tag}"
    hp = lambda tag: f"{{{HWPX_NS['hp']}}}{tag}"
    archive = zipfile.ZipFile(io.BytesIO(data))
    first = archive.infolist()[0]
    assert first.filename == "mimetype" and first.compress_type == zipfile.ZIP_STORED
    assert archive.read("mimetype") == b"application/hwp+zip"
    parts = {
        name: etree.fromstring(archive.read(name))
        for name in archive.namelist()
        if name.endswith((".xml", ".hpf"))
    }
    head = parts["Contents/header.xml"]
    for element in head.iter():
        if element.get("itemCnt") is not None:
            assert int(element.get("itemCnt")) == len(element), element.tag

    def ids(tag: str) -> set:
        values = [e.get("id") for e in head.iter(hh(tag))]
        assert len(values) == len(set(values)), f"duplicate {tag} id"
        return set(values)

    fonts = {}
    for face in head.iter(hh("fontface")):
        items = face.findall(hh("font"))
        assert int(face.get("fontCnt")) == len(items), face.get("lang")
        font_ids = [f.get("id") for f in items]
        assert len(font_ids) == len(set(font_ids))
        fonts[face.get("lang").lower()] = set(font_ids)
    chars, paras, fills, tabs, styles = (ids(t) for t in ("charPr", "paraPr", "borderFill", "tabPr", "style"))
    for char in head.iter(hh("charPr")):
        assert char.get("borderFillIDRef") in fills
        ref = char.find(hh("fontRef"))
        for lang, font_ids in fonts.items():
            assert ref.get(lang) in font_ids, (char.get("id"), lang)
    for para in head.iter(hh("paraPr")):
        assert para.get("tabPrIDRef") in tabs
        for border in para.iter(hh("border")):
            assert border.get("borderFillIDRef") in fills
    for style in head.iter(hh("style")):
        assert style.get("paraPrIDRef") in paras and style.get("charPrIDRef") in chars
        assert style.get("nextStyleIDRef") in styles
    sections = [n for n in parts if n.startswith("Contents/section")]
    if head.get("secCnt") is not None:
        assert int(head.get("secCnt")) == len(sections)
    section = parts["Contents/section0.xml"]
    for kind in ("header", "footer"):
        pages = [c.get("applyPageType") for c in section.iter(hp(kind))]
        assert len(pages) == len(set(pages)), f"two {kind}s for the same pages"
    for p in section.iter(hp("p")):
        assert p.get("paraPrIDRef") in paras and p.get("styleIDRef") in styles
    for run in section.iter(hp("run")):
        assert run.get("charPrIDRef") in chars
    for element in (*section.iter(hp("tbl")), *section.iter(hp("tc"))):
        assert element.get("borderFillIDRef") in fills, element.tag
    begins = {b.get("id") for b in section.iter(hp("fieldBegin"))}
    for end in section.iter(hp("fieldEnd")):
        assert end.get("beginIDRef") in begins
    sec_prs = list(section.iter(hp("secPr")))
    assert len(sec_prs) == 1
    assert sec_prs[0].getparent().getparent() is section[0]
    return parts


def hwpx_settings(**changes):
    """hwpx Settings with the built-in template's values, for writer tests below the adapter."""
    from dataclasses import replace as _replace

    from mdcast.hwpx.look import Look, Settings

    body = Look(font="맑은 고딕", size=11, bold=False, italic=False, color="#000000", align="justify",
                line_height=1.8, space_before=0, space_after=4, indent_left=0, keep_with_next=False)
    looks = {
        "body": body,
        "quote": _replace(body, size=10.5, color="#555555", align="left", space_after=0, indent_left=5),
        "code": Look(font="Consolas", font_hangul="맑은 고딕", size=9.5, bold=False, italic=False, color="#000000", align="left",
                     line_height=1.3, space_before=0, space_after=0, indent_left=0, background="#F6F8FA"),
        "caption": _replace(body, size=9.5, bold=True, align="center", line_height=1.3, space_after=0),
        "footnote": _replace(body, size=9, align="left", line_height=1.3, space_after=0),
    }
    for level in range(1, 7):
        looks[f"heading{level}"] = _replace(body, size=16 if level == 1 else 12, bold=True, align="left",
                                            line_height=1.3, space_before=12, space_after=6, keep_with_next=True)
    styles = {f"heading{level}": f"개요 {level}" for level in range(1, 7)}
    styles.update(body="바탕글", quote="바탕글", code="바탕글", caption="캡션", footnote="각주")
    markers = "□○-·"
    values = {
        "looks": looks, "styles": styles, "body_size": 11.0, "body_line_height": 1.8,
        "list_marker": lambda ordered, level, n: f"{n}." if ordered else markers[min(level, 4) - 1],
    }
    values.update(changes)
    return Settings(**values)


def hwpx_render(md: str, settings=None, doc=None):
    """Header, shapes, paragraphs and notices of ``md`` (or ``doc``) on the built-in container."""
    from mdcast.ast import md_to_doc
    from mdcast.hwpx.container import Package
    from mdcast.hwpx.header import Header
    from mdcast.hwpx.look import Notices
    from mdcast.hwpx.section import SectionWriter
    from mdcast.hwpx.shapes import build_shapes

    settings = settings or hwpx_settings()
    header = Header(Package.read(builtin_hwpx()).xml("Contents/header.xml"))
    notices = Notices()
    shapes = build_shapes(header, settings, notices)
    paragraphs = SectionWriter(header, shapes, settings, notices, width=170.0).build(doc or md_to_doc(md))
    return header, shapes, paragraphs, notices


def hwpx_texts(paragraphs) -> list[str]:
    """Each paragraph's own text (runs directly under it; header/footer lists excluded)."""
    t = f"{{{HWPX_NS['hp']}}}t"
    run = f"{{{HWPX_NS['hp']}}}run"
    return ["".join("".join(x.itertext()) for r in p.findall(run) for x in r.findall(t)) for p in paragraphs]


def _pdf_chars(data: bytes, page: int = 0):
    import ctypes

    import pypdfium2 as pdfium
    import pypdfium2.raw as pdfium_c

    pdf_page = pdfium.PdfDocument(data)[page]
    textpage = pdf_page.get_textpage()
    height = pdf_page.get_height()
    x, y = ctypes.c_double(), ctypes.c_double()
    chars = []
    for i in range(pdfium_c.FPDFText_CountChars(textpage.raw)):
        pdfium_c.FPDFText_GetCharOrigin(textpage.raw, i, ctypes.byref(x), ctypes.byref(y))
        ch = chr(pdfium_c.FPDFText_GetUnicode(textpage.raw, i) or 0xFFFD)
        chars.append((ch, x.value * 25.4 / 72, (height - y.value) * 25.4 / 72))
    return chars


def char_origin(data: bytes, needle: str, page: int = 0) -> tuple[float, float]:
    """Baseline origin (x, y from the top-left, mm) of the first character of ``needle``."""
    chars = _pdf_chars(data, page)
    text = "".join(ch for ch, _, _ in chars)
    _, x, y = chars[text.index(needle)]
    return x, y


def line_starts(data: bytes, needle: str, page: int = 0) -> list[tuple[float, float]]:
    """(x, y) mm of the first visible character of each line from ``needle`` on."""
    chars = _pdf_chars(data, page)
    text = "".join(ch for ch, _, _ in chars)
    starts: list[tuple[float, float]] = []
    for ch, x, y in chars[text.index(needle):]:
        if ch.strip() and (not starts or abs(y - starts[-1][1]) > 1.0):
            starts.append((x, y))
    return starts
