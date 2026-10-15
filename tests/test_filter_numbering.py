import panflute as pf

from mdcast.ast import md_to_doc
from mdcast.filters import numbering
from mdcast.filters.context import FilterContext
from mdcast.report import Report
from mdcast.templates.schema import deep_merge, load_builtin_default
from mdcast.types import Format


def run(md, overrides=None):
    config = deep_merge(load_builtin_default(), overrides or {})
    ctx = FilterContext(Format.DOCX, config, {}, Report("a.md", Format.MD))
    return numbering.apply(md_to_doc(md), ctx)


def headers(doc):
    found = []
    doc.walk(lambda e, d: found.append(pf.stringify(e)) if isinstance(e, pf.Header) else None)
    return found


def captions(doc):
    found = []

    def grab(e, d):
        if isinstance(e, (pf.Table, pf.Figure)) and e.caption is not None:
            found.append(pf.stringify(e.caption))

    doc.walk(grab)
    return found


def test_heading_numbers_reset_deeper_levels():
    doc = run("# 개요\n\n## 배경\n\n## 목적\n\n# 계획\n\n## 일정\n", {"headings": {"numbering": ["Ⅰ.", "1."]}})
    assert headers(doc) == ["Ⅰ. 개요", "1. 배경", "2. 목적", "Ⅱ. 계획", "1. 일정"]


def test_levels_beyond_patterns_and_unnumbered_are_skipped():
    doc = run("# A\n\n### deep\n\n# 부록 {.unnumbered}\n\n# B\n", {"headings": {"numbering": ["Ⅰ."]}})
    assert headers(doc) == ["Ⅰ. A", "deep", "부록", "Ⅱ. B"]


def test_no_patterns_means_no_change():
    assert headers(run("# A\n\n## B\n")) == ["A", "B"]


def test_table_and_figure_captions_numbered_independently():
    md = (
        "Table: 요약\n\n| a |\n|---|\n| 1 |\n\n"
        "![구성도](x.png)\n\n"
        "Table: 둘째\n\n| b |\n|---|\n| 2 |\n"
    )
    assert captions(run(md)) == ["<표 1> 요약", "[그림 1] 구성도", "<표 2> 둘째"]


def test_uncaptioned_table_is_not_numbered():
    # Controller ruling B1: pandoc attaches a "Table:" line to the table next to it,
    # so the uncaptioned table goes last to keep the captioned table's caption unambiguous.
    md = "Table: 캡션\n\n| a |\n|---|\n| 1 |\n\n| b |\n|---|\n| 2 |\n"
    assert captions(run(md)) == ["<표 1> 캡션", ""]


def number_attrs(doc):
    found = []

    def grab(e, d):
        if isinstance(e, pf.Header):
            found.append((pf.stringify(e), e.attributes.get("mdcast-number")))

    doc.walk(grab)
    return found


def test_numbered_headers_carry_their_visible_prefix():
    doc = run("# 개요\n\n## 배경\n\n## 목적\n", {"headings": {"numbering": ["1.", "가."]}})
    assert number_attrs(doc) == [("1. 개요", "1."), ("가. 배경", "가."), ("나. 목적", "나.")]


def test_unnumbered_and_too_deep_headers_get_no_number_attribute():
    doc = run("# A\n\n### deep\n\n# 부록 {.unnumbered}\n", {"headings": {"numbering": ["Ⅰ.", "1."]}})
    assert number_attrs(doc) == [("Ⅰ. A", "Ⅰ."), ("deep", None), ("부록", None)]


def test_number_attribute_restarts_with_the_deeper_counter():
    doc = run("# 하나\n\n## 가\n\n## 나\n\n# 둘\n\n## 가\n", {"headings": {"numbering": ["1.", "가."]}})
    assert [n for _, n in number_attrs(doc)] == ["1.", "가.", "나.", "2.", "가."]


def test_number_attribute_is_not_visible_in_docx_or_pdf(tmp_path):
    import io
    import zipfile

    import docx
    import pypdfium2 as pdfium
    from helpers import build, write_template

    from mdcast.writers.docx import write_docx
    from mdcast.writers.pdf import write_pdf

    tpl = write_template(tmp_path, "t", "name: t\nheadings: {numbering: ['1.', '가.']}\n")
    md = "# 개요\n\n## 배경\n\n본문\n"
    doc, ctx = build(tmp_path, md, Format.DOCX, template=str(tpl))
    data = write_docx(doc, ctx)
    assert "mdcast-number" not in zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode("utf-8")
    texts = [p.text for p in docx.Document(io.BytesIO(data)).paragraphs]
    assert texts[:2] == ["1. 개요", "가. 배경"]
    assert " ".join(texts).count("1.") == 1 and " ".join(texts).count("가.") == 1
    doc, ctx = build(tmp_path, md, Format.PDF, template=str(tpl))
    pdf = pdfium.PdfDocument(write_pdf(doc, ctx))
    text = " ".join(pdf[i].get_textpage().get_text_range() for i in range(len(pdf)))
    assert "mdcast-number" not in text
    assert text.count("1.") == 1 and text.count("가.") == 1
    assert "1. 개요" in " ".join(text.split()) and "가. 배경" in " ".join(text.split())
