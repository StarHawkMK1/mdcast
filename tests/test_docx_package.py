from __future__ import annotations

import io
import zipfile

import panflute as pf
import pypandoc
import pytest
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx_builder import add_page_break, add_para, docx_bytes, strip_mdcast_metadata
from helpers import export, tiny_png, write_template

from mdcast.ast import json_to_doc
from mdcast.readers.docx_package import PAGEBREAK_SENTINEL, prepare_docx
from mdcast.types import Format


def read_prepared(data: bytes) -> pf.Doc:
    return json_to_doc(pypandoc.convert_text(prepare_docx(data).data, "json", format="docx"))


def test_outline_level_paragraph_becomes_heading_for_pandoc():
    data = docx_bytes(lambda d: (add_para(d, "1. 큰 제목", bold=True, outline=0), add_para(d, "본문")))
    doc = read_prepared(data)
    assert isinstance(doc.content[0], pf.Header) and doc.content[0].level == 1


def _build_levels(d):
    style = d.styles.add_style("Outlined", 1)
    ppr = style.element.get_or_add_pPr()
    level = OxmlElement("w:outlineLvl")
    level.set(qn("w:val"), "1")
    ppr.append(level)
    add_para(d, "스타일 제목", style="Outlined")
    add_para(d, "본문 수준", outline=9)
    add_para(d, "", outline=0)


def test_style_outline_level_and_body_level_9():
    doc = read_prepared(docx_bytes(_build_levels))
    assert [type(b).__name__ for b in doc.content] == ["Header", "Para"]
    assert doc.content[0].level == 2


def _build_breaks(d):
    add_para(d, "앞")
    add_page_break(d)
    add_para(d, "쪽 나눔 문단", page_break_before=True)
    add_para(d, "뒤")


def test_page_breaks_become_sentinel_paragraphs():
    doc = read_prepared(docx_bytes(_build_breaks))
    texts = [pf.stringify(b).strip() for b in doc.content]
    assert texts.count(PAGEBREAK_SENTINEL) == 2
    assert "쪽 나눔 문단" in texts and "앞" in texts and "뒤" in texts


def test_page_break_inside_text_paragraph_keeps_text():
    def build(d):
        p = add_para(d, "앞 글")
        p.add_run().add_break(WD_BREAK.PAGE)
        p.add_run("뒤 글")

    doc = read_prepared(docx_bytes(build))
    texts = [pf.stringify(b).strip() for b in doc.content]
    assert texts[0] == PAGEBREAK_SENTINEL and "앞 글" in texts[1] and "뒤 글" in texts[1]


def build_with_title_author_header_and_comment(d):
    d.core_properties.title = "제목"
    d.core_properties.author = "작성자"
    first = add_para(d, "첫 문단")
    add_para(d, "")
    add_para(d, "둘째 문단")
    d.sections[0].header.paragraphs[0].text = "머리말 글"
    d.add_comment(first.runs, text="메모", author="a")


def test_source_text_and_facts():
    p = prepare_docx(docx_bytes(build_with_title_author_header_and_comment))
    assert p.source_text[:2] == ["첫 문단", "둘째 문단"]
    assert "머리말 글" not in p.source_text
    assert p.title == "제목" and p.author == "작성자"
    assert p.has_header_footer and p.has_comments and p.generator == {}


def test_source_text_excludes_sentinels_and_includes_notes():
    p = prepare_docx(docx_bytes(_build_breaks))
    assert p.source_text == ["앞", "쪽 나눔 문단", "뒤"]


def test_no_header_footer_or_comments_in_plain_docx():
    p = prepare_docx(docx_bytes(lambda d: add_para(d, "x")))
    assert not p.has_header_footer and not p.has_comments and p.table_styles == []


def test_image_only_header_counts_as_header_content():
    def build(d):
        add_para(d, "본문")
        d.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(tiny_png()))

    assert prepare_docx(docx_bytes(build)).has_header_footer


def test_table_variant_survives_in_table_styles(tmp_path):
    tpl = write_template(tmp_path, "kinds", """
        schema: 1
        name: kinds
        version: 1.0.0
        description: t
        extends: default
        tables:
          variants:
            plain: { header: { background: "#FFFFFF" } }
    """)
    md = "| a |\n|---|\n| 1 |\n\n<!-- table: plain -->\n\n| b |\n|---|\n| 2 |\n"
    assert prepare_docx(export(tmp_path, md, Format.DOCX, template=str(tpl))).table_styles == [None, "plain"]


def build_outer_nested_and_next_tables_with_styles(d):
    def style(table, name):
        if name is None:
            return
        created = d.styles.add_style(name, 3)
        pr = table._tbl.tblPr
        node = OxmlElement("w:tblStyle")
        node.set(qn("w:val"), created.style_id)
        pr.insert(0, node)

    outer = d.add_table(rows=1, cols=1)
    style(outer, "Table (a)")
    cell = outer.cell(0, 0)
    inner = cell.add_table(rows=1, cols=1)
    style(inner, "Table (b)")
    cell.add_paragraph("")
    d.add_paragraph("사이")
    d.add_table(rows=1, cols=1)


def test_table_styles_follow_document_order_with_nested_tables():
    data = docx_bytes(build_outer_nested_and_next_tables_with_styles)
    assert prepare_docx(data).table_styles == ["a", "b", None]


def test_generator_is_read_from_mdcast_docx(tmp_path):
    data = export(tmp_path, "# 개요\n\n본문\n", Format.DOCX)
    assert prepare_docx(data).generator["mdcast.template.name"] == "default"
    assert prepare_docx(strip_mdcast_metadata(data)).generator == {}


def test_other_entries_are_kept_byte_for_byte(tmp_path):
    data = export(tmp_path, "# 개요\n\n본문\n", Format.DOCX)
    before = zipfile.ZipFile(io.BytesIO(data))
    after = zipfile.ZipFile(io.BytesIO(prepare_docx(data).data))
    assert before.namelist() == after.namelist()
    for name in ("word/numbering.xml", "docProps/core.xml", "word/theme/theme1.xml"):
        assert before.read(name) == after.read(name)


def test_malformed_input_raises():
    with pytest.raises(zipfile.BadZipFile):
        prepare_docx(b"not a zip")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        z.writestr("x.txt", "x")
    with pytest.raises(KeyError):  # no word/document.xml
        prepare_docx(buffer.getvalue())
