import io
from pathlib import Path

import docx
import pypdfium2 as pdfium
import pytest
from helpers import HWPX_NS, hwpx_integrity

from mdcast import pipeline
from mdcast.report import Level
from mdcast.types import Format, Options, Source

REPO_ROOT = Path(__file__).resolve().parents[1]
FIXTURES = Path(__file__).parent / "fixtures" / "md"
TEMPLATE_WITH_PLAIN = (
    "name: e2e\n"
    "tables: {variants: {plain: {header: {background: null}}}}\n"
    "headings: {numbering: ['Ⅰ.', '1.']}\n"
)


@pytest.mark.parametrize("name", ["kitchen-sink.md", "korean-report.md", "edge-cases.md"])
def test_fixture_converts_to_docx_and_pdf(tmp_path, name):
    tpl = tmp_path / "e2e"
    tpl.mkdir()
    (tpl / "template.yaml").write_text(TEMPLATE_WITH_PLAIN, encoding="utf-8")
    path = FIXTURES / name
    source = Source(path.read_bytes(), name, Format.MD, base_dir=path.parent)
    result = pipeline.convert(
        source,
        [Format.DOCX, Format.PDF],
        str(tpl),
        Options(offline=True, cwd=tmp_path, home=tmp_path / "home"),
    )
    errors = [e for e in result.report.events if e.level is Level.ERROR]
    assert errors == []
    outputs = {o.format: o for o in result.outputs}
    document = docx.Document(io.BytesIO(outputs[Format.DOCX].data))
    assert any(p.text.startswith("Ⅰ.") for p in document.paragraphs)
    pdf = pdfium.PdfDocument(outputs[Format.PDF].data)
    assert len(pdf) >= 2
    if name == "kitchen-sink.md":
        assert len(document.tables) >= 3
        body = document.element.xml
        assert "<a:blip" in body
        assert 'w:type="page"' in body
    out_dir = REPO_ROOT / ".test-output" / "e2e"
    out_dir.mkdir(parents=True, exist_ok=True)
    for output in result.outputs:
        (out_dir / output.name).write_bytes(output.data)


SUBMISSION = Path(__file__).parent / "fixtures" / "templates" / "submission"


@pytest.mark.parametrize("name", ["tables-images.md", "kitchen-sink.md"])
def test_fixture_converts_to_hwpx_with_tables_pictures_and_notes(tmp_path, name):
    path = FIXTURES / name
    source = Source(path.read_bytes(), name, Format.MD, base_dir=path.parent)
    result = pipeline.convert(
        source,
        [Format.DOCX, Format.HWPX, Format.PDF],
        str(SUBMISSION),
        Options(offline=True, cwd=tmp_path, home=tmp_path / "home"),
    )
    assert [e for e in result.report.events if e.level is Level.ERROR] == []
    outputs = {o.format: o for o in result.outputs}
    parts = hwpx_integrity(outputs[Format.HWPX].data)
    section = parts["Contents/section0.xml"]
    if name == "tables-images.md":
        count = {tag: len(list(section.iter(f"{{{HWPX_NS['hp']}}}{tag}"))) for tag in ("tbl", "pic", "footNote")}
        assert count == {"tbl": 4, "pic": 5, "footNote": 3}
        assert not [e for e in result.report.events if e.where.get("format") == "hwpx" and e.level is not Level.INFO]
    out_dir = REPO_ROOT / ".test-output" / "e2e"
    out_dir.mkdir(parents=True, exist_ok=True)
    for output in result.outputs:
        (out_dir / output.name).write_bytes(output.data)


def test_edge_case_fixture_has_mixed_line_endings():
    data = (FIXTURES / "edge-cases.md").read_bytes()
    assert b"\r\n" in data
    assert b"\n" in data.replace(b"\r\n", b"")
