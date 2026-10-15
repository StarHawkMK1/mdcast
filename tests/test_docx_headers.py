import io

import docx
from docx.oxml.ns import qn
from helpers import build, make_reference_docx, write_template

from mdcast.fields import ResolvedFields
from mdcast.report import Codes, Report
from mdcast.types import Format
from mdcast.writers.docx import write_docx
from mdcast.writers.docx.headers import fill_placeholders


def render(tmp_path, md, template=None):
    doc, ctx = build(tmp_path, md, Format.DOCX, template=template)
    return docx.Document(io.BytesIO(write_docx(doc, ctx))), ctx


def xml_of(part) -> str:
    return part._element.xml


def test_yaml_slots_write_left_center_right(tmp_path):
    tpl = write_template(tmp_path, "t", "name: t\nheader: {left: '{{dept}}', center: '', right: '{{security}}'}\n")
    d, _ = render(tmp_path, "---\ndept: 전략기획팀\nsecurity: 대외비\n---\n\n본문\n", str(tpl))
    assert [p.text for p in d.sections[0].header.paragraphs] == ["전략기획팀\t\t대외비"]


def test_default_footer_has_page_fields_and_header_has_bottom_line(tmp_path):
    d, _ = render(tmp_path, "---\ntitle: 사업계획\n---\n\n본문\n")
    footer_xml = xml_of(d.sections[0].footer)
    assert 'w:instr=" PAGE "' in footer_xml and 'w:instr=" NUMPAGES "' in footer_xml
    header_xml = xml_of(d.sections[0].header)
    assert "<w:bottom " in header_xml and "사업계획" in header_xml


def test_reference_mode_fills_split_runs_and_page_field(tmp_path):
    tpl = write_template(tmp_path, "t", "name: t\nheader: {source: reference}\n")
    make_reference_docx(tpl / "reference.docx", header_runs=["부서: {{de", "pt}} / ", "{{page}}"])
    d, ctx = render(tmp_path, "---\ndept: 전략기획팀\n---\n\n본문\n", str(tpl))
    header = d.sections[0].header
    assert header.paragraphs[0].text == "부서: 전략기획팀 / "
    assert 'w:instr=" PAGE "' in xml_of(header)
    assert ctx.report.events == []


def fields(**values) -> ResolvedFields:
    return ResolvedFields(dict(values), Report("a.md", Format.MD))


def paragraph_with_runs(*texts):
    document = docx.Document()
    paragraph = document.add_paragraph()
    for text in texts:
        paragraph.add_run(text)
    return paragraph


def test_fill_two_placeholders_in_one_run_and_keep_surroundings():
    p = paragraph_with_runs("[{{a}}-{{b}}]")
    fill_placeholders(p, fields(a="가", b="나"))
    assert p.text == "[가-나]"


def test_fill_keeps_run_formatting_of_first_run():
    p = paragraph_with_runs("x{{ti", "tle}}y")
    p.runs[0].bold = True
    fill_placeholders(p, fields(title="제목"))
    assert p.text == "x제목y"
    assert p.runs[0].bold is True


def test_fill_missing_value_reports_field_empty():
    f = fields()
    p = paragraph_with_runs("{{dept}}")
    fill_placeholders(p, f)
    assert p.text == ""
    assert [e.code for e in f.report.events] == [Codes.FIELD_EMPTY]


def test_fill_page_field_mid_run():
    p = paragraph_with_runs("쪽 {{page}} 끝")
    fill_placeholders(p, fields())
    assert p.text == "쪽  끝"
    assert p._p.find(qn("w:fldSimple")) is not None


def _ref_with_first_page_header(path):
    make_reference_docx(path, header_runs=["{{dept}} 본문"])
    document = docx.Document(str(path))
    section = document.sections[0]
    section.different_first_page_header_footer = True
    first = section.first_page_header
    first.is_linked_to_previous = False
    first.paragraphs[0].text = "첫쪽 {{dept}}"
    document.save(str(path))


def test_mixed_reference_header_yaml_footer_keeps_first_page_header(tmp_path):
    tpl = write_template(
        tmp_path, "t", "name: t\nheader: {source: reference}\nfooter: {source: yaml, center: '{{page}}'}\n"
    )
    _ref_with_first_page_header(tpl / "reference.docx")
    d, _ = render(tmp_path, "---\ndept: 전략기획팀\n---\n\n본문\n", str(tpl))
    section = d.sections[0]
    assert section.different_first_page_header_footer is True
    assert section.first_page_header.paragraphs[0].text == "첫쪽 전략기획팀"
    assert section.header.paragraphs[0].text == "전략기획팀 본문"
    # the yaml footer must also show on the first page
    assert 'w:instr=" PAGE "' in xml_of(section.footer)
    assert 'w:instr=" PAGE "' in xml_of(section.first_page_footer)


def test_all_yaml_turns_title_page_off_and_replaces_even_parts(tmp_path):
    tpl = write_template(tmp_path, "t", "name: t\nheader: {left: 'L'}\nfooter: {center: 'C'}\n")
    ref = make_reference_docx(tpl / "reference.docx", header_runs=["x"])
    document = docx.Document(str(ref))
    document.settings.odd_and_even_pages_header_footer = True
    section = document.sections[0]
    section.different_first_page_header_footer = True
    section.even_page_header.is_linked_to_previous = False
    section.even_page_header.paragraphs[0].text = "짝수쪽 원본"
    document.save(str(ref))
    d, _ = render(tmp_path, "본문\n", str(tpl))
    section = d.sections[0]
    assert section.different_first_page_header_footer is False
    assert section.header.paragraphs[0].text.startswith("L")
    assert section.even_page_header.paragraphs[0].text.startswith("L")
    assert "짝수쪽" not in xml_of(section.even_page_header)
