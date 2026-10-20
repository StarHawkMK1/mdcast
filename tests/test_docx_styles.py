import io

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from helpers import build, make_reference_docx, write_template

from mdcast.report import Codes
from mdcast.types import Format
from mdcast.writers.docx import write_docx


def render(tmp_path, md, template=None):
    doc, ctx = build(tmp_path, md, Format.DOCX, template=template)
    return docx.Document(io.BytesIO(write_docx(doc, ctx))), ctx


def test_page_setup_from_template(tmp_path):
    tpl = write_template(tmp_path, "t", "name: t\npage: {size: A5, orientation: landscape, margin: {top: 10mm, bottom: 10mm, left: 12mm, right: 12mm}}\n")
    d, _ = render(tmp_path, "본문\n", str(tpl))
    s = d.sections[0]
    assert (round(s.page_width.mm), round(s.page_height.mm)) == (210, 148)
    assert round(s.top_margin.mm) == 10 and round(s.left_margin.mm) == 12


def test_default_element_looks(tmp_path):
    d, _ = render(tmp_path, "# 제목\n\n본문\n")
    h1 = d.styles["Heading 1"]
    assert h1.font.size.pt == 16 and str(h1.font.color.rgb) == "1F3864"
    body = d.styles["Body Text"]
    assert body.font.size.pt == 11
    assert body.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert _spacing(body) == ("396", "exact")
    rfonts = d.styles["Normal"].element.rPr.rFonts
    assert rfonts.get(qn("w:eastAsia")) == "맑은 고딕"
    assert rfonts.get(qn("w:ascii")) == "Malgun Gothic"
    assert rfonts.get(qn("w:asciiTheme")) is None


def test_heading_paragraph_remapped_to_custom_style(tmp_path):
    tpl = write_template(tmp_path, "t", "name: t\nstyles: {heading1: {docx: 보고서 제목1}}\nelements: {heading1: {size: 16pt}}\n")
    make_reference_docx(tpl / "reference.docx", extra_styles=("보고서 제목1",))
    d, ctx = render(tmp_path, "# 제목\n\n본문\n", str(tpl))
    heading = next(p for p in d.paragraphs if p.text == "제목")
    assert heading.style.name == "보고서 제목1"
    assert d.styles["보고서 제목1"].font.size.pt == 16
    assert not [e for e in ctx.report.events if e.code == Codes.STYLE_UNMAPPED]


def test_missing_mapped_style_is_info_and_keeps_pandoc_style(tmp_path):
    tpl = write_template(tmp_path, "t", "name: t\nstyles: {heading1: {docx: 없는스타일}}\n")
    d, ctx = render(tmp_path, "# 제목\n", str(tpl))
    heading = next(p for p in d.paragraphs if p.text == "제목")
    assert heading.style.name == "Heading 1"
    assert [(e.code, e.where.get("format")) for e in ctx.report.events if e.code == Codes.STYLE_UNMAPPED] == [
        (Codes.STYLE_UNMAPPED, "docx")
    ]


def test_code_block_style_uses_code_font(tmp_path):
    d, _ = render(tmp_path, "```\nprint(1)\n```\n")
    names = {s.name for s in d.styles}
    assert "Source Code" in names
    assert d.styles["Source Code"].font.name == "Consolas"


def test_corrupt_custom_reference_falls_back_to_builtin_container(tmp_path):
    tpl = write_template(tmp_path, "t", "name: t\n", {"reference.docx": b"not a zip"})
    d, ctx = render(tmp_path, "# 제목\n", str(tpl))
    assert any(p.text == "제목" for p in d.paragraphs)
    assert [(e.code, e.where.get("format")) for e in ctx.report.events if e.code == Codes.TEMPLATE_INVALID] == [
        (Codes.TEMPLATE_INVALID, "docx")
    ]


def test_custom_container_keeps_its_own_values_unless_yaml_sets_them(tmp_path):
    tpl = write_template(tmp_path, "t", "name: t\nelements: {body: {size: 13pt}}\n")
    ref = make_reference_docx(tpl / "reference.docx")
    reference = docx.Document(str(ref))
    reference.styles["Heading 1"].font.size = Pt(20)
    reference.sections[0].page_width = Mm(150)
    reference.save(str(ref))
    d, _ = render(tmp_path, "# 제목\n\n본문\n", str(tpl))
    assert d.styles["Heading 1"].font.size.pt == 20
    assert d.styles["Body Text"].font.size.pt == 13
    assert round(d.sections[0].page_width.mm) == 150


def styles_by_name(document):
    """Styles keyed by their UI name; ``styles["Caption"]`` would fall back to a style-id lookup."""
    return {s.name: s for s in document.styles}


def company_reference(tpl, *, caption_base=None):
    """A reference.docx with company paragraph styles 본문 (13pt, own font) and 캡션."""
    from docx.enum.style import WD_STYLE_TYPE

    ref = make_reference_docx(tpl / "reference.docx")
    reference = docx.Document(str(ref))
    body = reference.styles.add_style("본문", WD_STYLE_TYPE.PARAGRAPH)
    body.font.size = Pt(13)
    body.font.name = "Batang"
    caption = reference.styles.add_style("캡션", WD_STYLE_TYPE.PARAGRAPH)
    if caption_base:
        caption.base_style = styles_by_name(reference)[caption_base]
    reference.save(str(ref))
    return ref


def test_secondary_body_styles_follow_the_mapped_body_style(tmp_path):
    tpl = write_template(tmp_path, "t", "name: t\nstyles: {body: {docx: 본문}, caption: {docx: 캡션}}\n")
    company_reference(tpl)
    d, ctx = render(tmp_path, "# 제목\n\n첫 문단\n\n둘째 문단\n\n- 하나\n- 둘\n\n![그림](none.png)\n", str(tpl))
    styles = styles_by_name(d)
    body_id = styles["본문"].style_id
    by_text = {p.text: p.style.name for p in d.paragraphs}
    assert by_text["첫 문단"] == "First Paragraph" and by_text["하나"] == "Compact"
    assert by_text["둘째 문단"] == "본문"
    for name in ("First Paragraph", "Compact"):
        assert styles[name].element.basedOn_val == body_id, name
    assert styles["First Paragraph"].font.size is None  # inherits 본문's 13pt
    assert styles["Compact"].paragraph_format.space_after is not None  # keeps its tight spacing
    caption_id = styles["캡션"].style_id
    for name in ("Image Caption", "Caption"):
        assert styles[name].element.basedOn_val == caption_id, name
    assert not [e for e in ctx.report.events if e.code == Codes.STYLE_UNMAPPED]


def test_secondary_style_the_target_is_based_on_is_not_rebased(tmp_path):
    tpl = write_template(tmp_path, "t", "name: t\nstyles: {caption: {docx: 캡션}}\n")
    company_reference(tpl, caption_base="Caption")
    d, _ = render(tmp_path, "![그림](none.png)\n", str(tpl))
    styles = styles_by_name(d)
    company_id = styles["캡션"].style_id
    assert styles["캡션"].element.basedOn_val == styles["Caption"].style_id
    assert styles["Caption"].element.basedOn_val != company_id
    assert styles["Image Caption"].element.basedOn_val == company_id


def _autospace(d) -> dict[str, str | None]:
    ppr = d.styles.element.find(qn("w:docDefaults")).find(qn("w:pPrDefault")).find(qn("w:pPr"))
    return {tag: (el.get(qn("w:val")) if (el := ppr.find(qn(f"w:{tag}"))) is not None else None) for tag in ("autoSpaceDE", "autoSpaceDN")}


def test_builtin_container_turns_off_korean_digit_autospace(tmp_path):
    # Word widens "2026년" to "2026 년" by default; Typst does not, so the PDF would differ
    d, _ = render(tmp_path, "2026년 4분기\n")
    assert _autospace(d) == {"autoSpaceDE": "0", "autoSpaceDN": "0"}


def test_custom_container_keeps_its_own_autospace(tmp_path):
    tpl = write_template(tmp_path, "t", "name: t\n")
    make_reference_docx(tpl / "reference.docx")
    d, _ = render(tmp_path, "2026년 4분기\n", str(tpl))
    assert _autospace(d) == {"autoSpaceDE": None, "autoSpaceDN": None}


def test_default_captions_are_not_italic(tmp_path):
    # pandoc's reference makes Caption italic; the PDF caption is upright
    d, _ = render(tmp_path, "Table: 요약\n\n| a |\n|---|\n| 1 |\n")
    styles = styles_by_name(d)
    for name in ("Table Caption", "Caption"):
        assert styles[name].font.italic is False


def _spacing(style) -> tuple[str | None, str | None]:
    ppr = style.element.pPr
    spacing = None if ppr is None else ppr.find(qn("w:spacing"))
    return (None, None) if spacing is None else (spacing.get(qn("w:line")), spacing.get(qn("w:lineRule")))


def test_builtin_styles_use_exact_line_height(tmp_path):
    d, _ = render(tmp_path, "# 제목\n\n본문\n\nTable: 캡션\n\n| a |\n|---|\n| 1 |\n")
    styles = styles_by_name(d)
    assert _spacing(styles["Body Text"]) == ("396", "exact")  # 11pt × 180% × 20
    assert _spacing(styles["First Paragraph"]) == ("396", "exact")
    assert _spacing(styles["Heading 1"]) == ("416", "exact")  # 16pt × 130%
    assert _spacing(styles["Table Caption"]) == ("247", "exact")  # 9.5pt × 130%
    assert _spacing(styles["Footnote Text"]) == ("234", "exact")  # 9pt × 130%


def test_builtin_compact_has_no_gaps_and_inherits_the_body_line(tmp_path):
    d, _ = render(tmp_path, "- 하나\n- 둘\n")
    compact = styles_by_name(d)["Compact"]
    assert compact.paragraph_format.space_before == 0 and compact.paragraph_format.space_after == 0
    assert _spacing(compact)[0] is None  # line height comes from Body Text


def _size_pt(d, style) -> float:
    """Effective font size: the style, its bases, then the document defaults (as the writer does)."""
    while style is not None:
        if style.font.size is not None:
            return style.font.size.pt
        style = style.base_style
    sz = d.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}/{qn('w:sz')}")
    return int(sz.get(qn("w:val"))) / 2


def test_custom_container_sets_only_the_written_line_heights(tmp_path):
    tpl = write_template(tmp_path, "t", "name: t\nelements: {body: {line_height: 150%, size: 12pt}, heading1: {line_height: 120%}}\n")
    make_reference_docx(tpl / "reference.docx")
    d, _ = render(tmp_path, "# 제목\n\n본문\n", str(tpl))
    styles = styles_by_name(d)
    assert _spacing(styles["Body Text"]) == ("360", "exact")  # 12pt × 150%
    h1 = styles["Heading 1"]
    assert _spacing(h1) == (str(round(1.2 * _size_pt(d, h1) * 20)), "exact")  # the reference's own size
    assert _spacing(styles["Heading 2"])[1] != "exact"  # not written: the reference's own
    assert styles["Compact"].paragraph_format.space_after != 0  # tightening is built-in only


def test_builtin_code_runs_take_the_code_size(tmp_path):
    # pandoc's Verbatim Char fixes 11pt on every code run; the code block's 9.5pt must show
    d, _ = render(tmp_path, "```\ncode\n```\n\n본문 `inline`\n")
    rpr = styles_by_name(d)["Verbatim Char"].element.rPr
    assert rpr is None or (rpr.find(qn("w:sz")) is None and rpr.find(qn("w:szCs")) is None)


def test_line_height_stays_within_words_range():
    from mdcast.writers.docx.xml import line_twips

    assert line_twips(20, 1000) == 31680  # Word's maximum, 1584pt
    assert line_twips(0.001, 1) == 1
    assert line_twips(1.8, 11) == 396


def _note_paragraphs(d) -> dict[str, object]:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from lxml import etree

    part = next(r.target_part for r in d.part.rels.values() if r.reltype == RT.FOOTNOTES)
    root = etree.fromstring(part.blob)
    return {"".join(p.itertext()).strip(): p for p in root.iter(qn("w:p"))}


def test_tight_list_and_cells_in_a_footnote_take_the_footnote_look(tmp_path):
    # pandoc gives them Compact (the body's size and line); the PDF draws them as footnote text
    md = "본문[^1]\n\n[^1]: 각주\n\n    - 하나\n    - 둘\n\n    | 칸 |\n    |---|\n    | 셀 |\n"
    d, _ = render(tmp_path, md)
    found = _note_paragraphs(d)
    for text in ("하나", "둘", "칸", "셀"):
        ppr = found[text].find(qn("w:pPr"))
        assert ppr.find(qn("w:pStyle")).get(qn("w:val")) == "FootnoteText", text
        spacing = ppr.find(qn("w:spacing"))
        assert (spacing.get(qn("w:before")), spacing.get(qn("w:after"))) == ("0", "0"), text
