from __future__ import annotations

import io
import zipfile
from collections.abc import Callable

import docx
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

_CUSTOM = "docProps/custom.xml"
_CP = "{http://schemas.openxmlformats.org/officeDocument/2006/custom-properties}"


def docx_bytes(build: Callable[[docx.document.Document], None]) -> bytes:
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def add_para(
    d,
    text: str,
    *,
    bold: bool = False,
    outline: int | None = None,
    page_break_before: bool = False,
    style: str | None = None,
) -> Paragraph:
    paragraph = d.add_paragraph(style=style)
    if text:
        paragraph.add_run(text).bold = bold or None
    if page_break_before:
        paragraph.paragraph_format.page_break_before = True
    if outline is not None:
        outline_level = OxmlElement("w:outlineLvl")
        outline_level.set(qn("w:val"), str(outline))
        paragraph._p.get_or_add_pPr().append(outline_level)
    return paragraph


def add_page_break(d) -> None:
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def strip_mdcast_metadata(data: bytes) -> bytes:
    """Drop the ``mdcast.*`` custom properties and the core ``comments``: an mdcast docx as someone else's."""
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            blob = source.read(info.filename)
            if info.filename == _CUSTOM:
                root = etree.fromstring(blob)
                for prop in root.findall(f"{_CP}property"):
                    if (prop.get("name") or "").startswith("mdcast."):
                        root.remove(prop)
                blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif info.filename == "docProps/core.xml":
                root = etree.fromstring(blob)
                for node in root.findall("{http://purl.org/dc/elements/1.1/}description"):
                    root.remove(node)
                blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(info, blob)
    return out.getvalue()
