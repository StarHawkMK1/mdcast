import http.server
import io
import threading
import zipfile
from contextlib import contextmanager
from dataclasses import replace

import docx
import panflute as pf
import pypdfium2 as pdfium
import pytest
from helpers import build, make_reference_docx, tiny_bmp, tiny_jpeg, tiny_png

from mdcast import pipeline
from mdcast.report import Codes
from mdcast.types import Format, Options, Source
from mdcast.writers.images import image_info, localize_images, webp_to_png


def images(doc):
    found = []
    doc.walk(lambda e, d: found.append(e.url) if isinstance(e, pf.Image) else None)
    return found


def test_relative_image_with_korean_and_space_is_copied(tmp_path):
    (tmp_path / "그림 1.png").write_bytes(tiny_png())
    doc, ctx = build(tmp_path, "![설명](<그림 1.png>)\n", Format.DOCX)
    doc = localize_images(doc, ctx)
    assert images(doc) == ["assets/img1.png"]
    assert (ctx.work_dir / "assets" / "img1.png").read_bytes() == tiny_png()
    assert ctx.report.events == []


def test_missing_image_becomes_placeholder_with_warning(tmp_path):
    doc, ctx = build(tmp_path, "앞 ![없는 그림](none.png) 뒤\n", Format.PDF)
    doc = localize_images(doc, ctx)
    assert images(doc) == []
    assert "[이미지: 없는 그림]" in pf.stringify(doc)
    assert [(e.code, e.where["format"]) for e in ctx.report.events] == [(Codes.IMAGE_FETCH_FAILED, "pdf")]


def test_remote_image_offline_is_not_fetched(tmp_path):
    doc, ctx = build(tmp_path, "![로고](https://example.com/logo.png)\n", Format.DOCX, offline=True)
    doc = localize_images(doc, ctx)
    assert images(doc) == []
    assert [e.code for e in ctx.report.events] == [Codes.IMAGE_FETCH_FAILED]


def test_absolute_path_image(tmp_path):
    target = tmp_path / "abs.png"
    target.write_bytes(tiny_png())
    doc, ctx = build(tmp_path, f"![a](<{target.as_posix()}>)\n", Format.DOCX)
    doc = localize_images(doc, ctx)
    assert images(doc) == ["assets/img1.png"]


def test_remote_image_unreachable_online_becomes_placeholder(tmp_path):
    doc, ctx = build(tmp_path, "![x](http://127.0.0.1:9/a.png)\n", Format.DOCX, offline=False)
    ctx.options = replace(ctx.options, url_timeout_s=0.5)
    doc = localize_images(doc, ctx)
    assert images(doc) == []
    assert [e.code for e in ctx.report.events] == [Codes.IMAGE_FETCH_FAILED]


def test_make_reference_docx_adds_styles_and_footer(tmp_path):
    path = make_reference_docx(tmp_path / "ref.docx", footer_runs=["a", "b"], extra_styles=("MyStyle",))
    document = docx.Document(str(path))
    assert "MyStyle" in [s.name for s in document.styles]
    assert document.sections[0].footer.paragraphs[0].text == "ab"


@contextmanager
def serve(handler_body):
    class Handler(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            handler_body(self)

        def log_message(self, *args):
            pass

    server = http.server.HTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}"
    finally:
        server.shutdown()
        server.server_close()


def failed_codes(tmp_path, url):
    doc, ctx = build(tmp_path, f"![x]({url})\n", Format.DOCX, offline=False)
    ctx.options = replace(ctx.options, url_timeout_s=2.0)
    doc = localize_images(doc, ctx)
    return images(doc), [e.code for e in ctx.report.events]


def test_directory_as_image_path_becomes_placeholder(tmp_path):
    (tmp_path / "dir.png").mkdir()
    doc, ctx = build(tmp_path, "![d](dir.png)\n", Format.DOCX)
    doc = localize_images(doc, ctx)
    assert images(doc) == []
    assert [e.code for e in ctx.report.events] == [Codes.IMAGE_FETCH_FAILED]


def test_unreadable_local_file_becomes_placeholder(tmp_path, monkeypatch):
    (tmp_path / "a.png").write_bytes(tiny_png())

    def boom(self):
        raise PermissionError("denied")

    monkeypatch.setattr("pathlib.Path.read_bytes", boom)
    doc, ctx = build(tmp_path, "![d](a.png)\n", Format.DOCX)
    doc = localize_images(doc, ctx)
    assert images(doc) == []
    assert [e.code for e in ctx.report.events] == [Codes.IMAGE_FETCH_FAILED]


def test_short_body_response_becomes_placeholder(tmp_path):
    def body(h):
        h.send_response(200)
        h.send_header("Content-Length", "1000")
        h.end_headers()
        h.wfile.write(b"abc")
        h.wfile.flush()
        h.close_connection = True

    with serve(body) as base:
        assert failed_codes(tmp_path, base + "/a.png") == ([], [Codes.IMAGE_FETCH_FAILED])


def test_oversized_image_becomes_placeholder(tmp_path, monkeypatch):
    monkeypatch.setattr("mdcast.writers.images.MAX_IMAGE_BYTES", 10)

    def body(h):
        h.send_response(200)
        h.send_header("Content-Length", str(len(tiny_png())))
        h.end_headers()
        h.wfile.write(tiny_png())

    with serve(body) as base:
        assert failed_codes(tmp_path, base + "/a.png") == ([], [Codes.IMAGE_FETCH_FAILED])


def test_redirect_to_non_http_is_refused(tmp_path):
    def body(h):
        h.send_response(302)
        h.send_header("Location", "ftp://127.0.0.1/a.png")
        h.end_headers()

    with serve(body) as base:
        assert failed_codes(tmp_path, base + "/a.png") == ([], [Codes.IMAGE_FETCH_FAILED])


def test_reachable_remote_image_is_downloaded(tmp_path):
    def body(h):
        h.send_response(200)
        h.send_header("Content-Length", str(len(tiny_png())))
        h.end_headers()
        h.wfile.write(tiny_png())

    with serve(body) as base:
        assert failed_codes(tmp_path, base + "/a.png") == (["assets/img1.png"], [])


def localized(tmp_path, name, data, md=None, *, strict_types=False):
    (tmp_path / name).write_bytes(data)
    target = Format.PDF if strict_types else Format.DOCX
    doc, ctx = build(tmp_path, md or f"앞 ![그림]({name}) 뒤\n", target)
    return localize_images(doc, ctx, strict_types=strict_types), ctx


def test_png_bytes_named_jpg_are_saved_as_png(tmp_path):
    doc, ctx = localized(tmp_path, "photo.jpg", tiny_png())
    assert images(doc) == ["assets/img1.png"]
    assert ctx.report.events == []


def test_extensionless_jpeg_is_saved_as_jpg(tmp_path):
    doc, ctx = localized(tmp_path, "photo", tiny_jpeg())
    assert images(doc) == ["assets/img1.jpg"]
    assert (ctx.work_dir / "assets" / "img1.jpg").read_bytes() == tiny_jpeg()


def test_sniffs_gif_webp_and_svg(tmp_path):
    svg = b'\xef\xbb\xbf<?xml version="1.0"?>\n<svg xmlns="http://www.w3.org/2000/svg" width="1" height="1"/>'
    (tmp_path / "a.bin").write_bytes(b"GIF89a" + bytes(20))
    (tmp_path / "b.bin").write_bytes(b"RIFF\x10\x00\x00\x00WEBPVP8 " + bytes(8))
    doc, _ = localized(tmp_path, "c.bin", svg, "![a](a.bin) ![b](b.bin) ![c](c.bin)\n")
    assert images(doc) == ["assets/img1.gif", "assets/img2.webp", "assets/img3.svg"]


def test_svg_after_declaration_comment_and_doctype_is_accepted(tmp_path):
    svg = (
        b'<?xml version="1.0" encoding="UTF-8"?>\n<!-- made by a tool -->\n'
        b'<!DOCTYPE svg PUBLIC "-//W3C//DTD SVG 1.1//EN" "http://www.w3.org/Graphics/SVG/1.1/DTD/svg11.dtd">\n'
        b'<SVG xmlns="http://www.w3.org/2000/svg" width="1" height="1"></SVG>'
    )
    doc, ctx = localized(tmp_path, "a.svg", svg)
    assert images(doc) == ["assets/img1.svg"] and ctx.report.events == []


def corrupt_png() -> bytes:
    data = bytearray(tiny_png())
    data[20] ^= 0xFF  # inside IHDR: its CRC no longer matches
    return bytes(data)


@pytest.mark.parametrize(
    ("name", "data", "reason"),
    [
        ("pic.bmp", tiny_bmp(), "지원하지 않는 이미지 형식"),
        ("cut.png", tiny_png()[:30], "손상"),
        ("crc.png", corrupt_png(), "손상"),
        ("empty.png", b"", "지원하지 않는 이미지 형식"),
        ("page.svg", b"<!DOCTYPE html>\n<html><body><svg width='1' height='1'></svg></body></html>", "지원하지 않는 이미지 형식"),
        ("icon.svg", b"<div><svg/></div>", "지원하지 않는 이미지 형식"),
    ],
)
def test_strict_types_turn_unsupported_or_broken_image_into_placeholder(tmp_path, name, data, reason):
    doc, ctx = localized(tmp_path, name, data, strict_types=True)
    assert images(doc) == []
    assert "[이미지: 그림]" in pf.stringify(doc)
    [event] = ctx.report.events
    assert event.code == Codes.IMAGE_FETCH_FAILED and reason in event.message


def test_remote_extensionless_jpeg_is_saved_as_jpg(tmp_path):
    def body(h):
        h.send_response(200)
        h.send_header("Content-Length", str(len(tiny_jpeg())))
        h.end_headers()
        h.wfile.write(tiny_jpeg())

    with serve(body) as base:
        assert failed_codes(tmp_path, base + "/photo") == (["assets/img1.jpg"], [])


def figures(doc):
    found = []
    doc.walk(lambda e, d: found.append(e) if isinstance(e, pf.Figure) else None)
    return found


@pytest.mark.parametrize("url", ["none.png", "data:image/png;base64,iVBORw0KGgo="])
def test_failed_figure_image_replaces_whole_figure_keeping_caption(tmp_path, url):
    doc, ctx = build(tmp_path, f"앞 문단\n\n![그림 설명]({url})\n\n뒤 문단\n", Format.DOCX)
    [figure] = figures(doc)
    caption = pf.stringify(figure.caption).strip()
    doc = localize_images(doc, ctx)
    assert figures(doc) == []
    assert [type(b) for b in doc.content] == [pf.Para, pf.Para, pf.Para]
    text = pf.stringify(doc.content[1])
    assert "[이미지: 그림 설명]" in text and caption in text
    assert [e.code for e in ctx.report.events] == [Codes.IMAGE_FETCH_FAILED]


def test_working_figure_is_kept(tmp_path):
    (tmp_path / "a.png").write_bytes(tiny_png())
    doc, ctx = build(tmp_path, "![그림 설명](a.png)\n", Format.DOCX)
    doc = localize_images(doc, ctx)
    assert len(figures(doc)) == 1 and images(doc) == ["assets/img1.png"]


def test_check_hook_rejection_becomes_placeholder(tmp_path):
    (tmp_path / "a.png").write_bytes(tiny_png())
    doc, ctx = build(tmp_path, "앞 ![그림](a.png) 뒤\n", Format.PDF)
    doc = localize_images(doc, ctx, check=lambda path: f"읽을 수 없음 {path}")
    assert images(doc) == []
    [event] = ctx.report.events
    assert event.code == Codes.IMAGE_FETCH_FAILED and "읽을 수 없음 assets/img1.png" in event.message


@pytest.mark.parametrize(
    ("name", "data", "expected"),
    [
        ("pic.bmp", tiny_bmp(), "assets/img1.bmp"),
        ("crc.png", corrupt_png(), "assets/img1.png"),
        ("drawing.emf", b"\x01\x00\x00\x00" + bytes(80), "assets/img1.emf"),
        ("noext", tiny_bmp(), "assets/img1.bmp"),  # BMP is recognised by its bytes (M2b spec §5.1)
    ],
)
def test_docx_keeps_types_it_does_not_recognise(tmp_path, name, data, expected):
    doc, ctx = localized(tmp_path, name, data)
    assert images(doc) == [expected]
    assert ctx.report.events == []


def test_bmp_is_embedded_in_docx_and_a_placeholder_in_pdf(tmp_path):
    (tmp_path / "pic.bmp").write_bytes(tiny_bmp())
    source = Source("앞 ![작은 그림](pic.bmp) 뒤\n".encode(), "doc.md", Format.MD, base_dir=tmp_path)
    options = Options(offline=True, cwd=tmp_path, home=tmp_path / "home")
    result = pipeline.convert(source, [Format.DOCX, Format.PDF], options=options)
    outputs = {o.format: o.data for o in result.outputs}
    names = zipfile.ZipFile(io.BytesIO(outputs[Format.DOCX])).namelist()
    media = [n for n in names if n.startswith("word/media/")]
    assert media and media[0].endswith(".bmp")
    pdf = pdfium.PdfDocument(outputs[Format.PDF])
    assert "[이미지: 작은 그림]" in pdf[0].get_textpage().get_text_range()
    failed = [e.where.get("format") for e in result.report.events if e.code == Codes.IMAGE_FETCH_FAILED]
    assert failed == ["pdf"]


def test_image_info_reads_sizes_and_dpi(tmp_path):
    from PIL import Image as PILImage

    cases = {}
    for name, kwargs in {"a.png": {"dpi": (300, 300)}, "b.png": {}, "c.jpg": {"dpi": (72, 72)},
                         "d.gif": {}, "e.bmp": {}, "f.webp": {}}.items():
        path = tmp_path / name
        PILImage.new("RGB", (40, 20), (1, 2, 3)).save(path, **kwargs)
        cases[name] = image_info(path.read_bytes())
    assert cases["a.png"] == ("png", 40, 20, 300.0)
    assert cases["b.png"] == ("png", 40, 20, 96.0)
    assert cases["c.jpg"] == ("jpg", 40, 20, 72.0)
    assert cases["d.gif"][:3] == ("gif", 40, 20)
    assert cases["e.bmp"][:3] == ("bmp", 40, 20)
    assert cases["f.webp"][:3] == ("webp", 40, 20)


@pytest.mark.parametrize("svg, size", [
    ('<svg xmlns="http://www.w3.org/2000/svg" width="400" height="200"/>', (400, 200)),
    ('<svg xmlns="http://www.w3.org/2000/svg" width="72pt" height="1in"/>', (96, 96)),
    ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 120 60"/>', (120, 60)),
    ('<svg xmlns="http://www.w3.org/2000/svg"/>', (300, 150)),
])
def test_image_info_reads_svg_sizes(svg, size):
    assert image_info(svg.encode()) == ("svg", *size, 96.0)


def test_image_info_rejects_unknown_bytes():
    with pytest.raises(ValueError):
        image_info(b"not an image")


def test_webp_becomes_a_png_of_the_same_pixels(tmp_path):
    from PIL import Image as PILImage

    path = tmp_path / "x.webp"
    PILImage.new("RGB", (40, 20), (1, 2, 3)).save(path)
    png = webp_to_png(path.read_bytes(), 40, tmp_path)
    assert image_info(png)[:3] == ("png", 40, 20)
